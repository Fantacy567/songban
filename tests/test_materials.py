"""SQLite/API tests. fixture_model is explicitly a protocol test, never a model fallback."""
import base64
import copy
import io
import json
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from fastapi.testclient import TestClient
from app.main import create_app
from app.materials_parse import parse_source
from app.planner import PlannerError
from tests import test_team_core as core_fixture


def fixture_model(context):
    items = []
    for source in context['sources']:
        for seg in source['segments']:
            for line in seg['text'].splitlines():
                if not line.strip():
                    continue
                category = '同伴想法' if '愿意' in line else '时间信息' if '日期' in line else '资料要求'
                items.append({'category': category, 'text': line, 'kind': 'statement', 'importance': 'high', 'citations': [{'source_id': source['id'], 'segment_id': seg['id'], 'quote': line}], 'member_ids': [], 'conflict': '人数冲突' if '人组队' in line else ''})
    items.append({'category': '待确认事项', 'text': '正式提交截止时间尚未提供，请向主办方核实。', 'kind': 'unknown', 'importance': 'normal', 'citations': [], 'member_ids': [], 'conflict': ''})
    return {'summary': '协议测试结果：对照来源核对要求与同伴想法。', 'items': items[:60]}


class MaterialsTests(unittest.TestCase):
    def setUp(self):
        core_fixture.CoreTests.setUp(self)
        self.client.close()
        self.app = create_app(self.path, core_only=True, materials_model_runner=fixture_model)
        self.client = TestClient(self.app)
        self.space = self.client.get('/api/materials/space', headers=self.h('a')).json()['id']
        self.url = '/api/materials/spaces/' + self.space
        self.counter = 0

    tearDown = core_fixture.CoreTests.tearDown

    def snap(self):
        return self.client.get(self.url, headers=self.h('a')).json()

    def key(self):
        self.counter += 1
        return 'request-' + str(self.counter)

    def add(self, text='活动日期：11月7日至8日。\n支持1至3人组队。\n小李愿意做设计。', **extra):
        body = {'request_id': self.key(), 'kind': 'text', 'name': '规则与聊天', 'content': text, **extra}
        r = self.client.post(self.url + '/sources', headers=self.h('a'), json=body)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()['source']

    def read(self, source):
        r = self.client.post(self.url + '/sources/' + source['id'] + '/parse', headers=self.h('a'), json={'expected_revision': self.snap()['revision']})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def generate(self, client=None, **extra):
        return (client or self.client).post(self.url + '/generate', headers=self.h('a'), json={'request_id': self.key(), 'expected_revision': self.snap()['revision'], **extra})

    def prepared(self):
        self.read(self.add(nature='sample'))
        r = self.generate(); self.assertEqual(r.status_code, 200, r.text)
        return self.snap()

    def test_dynamic_categories_and_real_locatable_citations(self):
        s = self.prepared()
        self.assertGreaterEqual(len({i['category'] for i in s['result']['items']}), 4)
        for item in s['result']['items']:
            for c in item['citations']:
                source = next(r for r in s['sources'] if r['id'] == c['source_id'])
                segment = next(g for g in source['segments'] if g['id'] == c['segment_id'])
                self.assertEqual(segment['text'][c['start']:c['end']], c['quote'])
                self.assertEqual(c['nature'], 'sample')
        self.assertIn('协议测试', s['result']['provenance'])
        self.assertFalse(s['stale'])

    def test_repeated_upload_and_request_payload_mismatch(self):
        a = self.add(request_id='same-request'); b = self.add(request_id='same-request')
        self.assertEqual(a['id'], b['id']); self.assertEqual(len(self.snap()['sources']), 1)
        r = self.client.post(self.url+'/sources', headers=self.h('a'), json={'kind':'text','name':'规则与聊天','content':'different','request_id':'same-request'})
        self.assertEqual(r.status_code,409)

    def test_repeated_generation_calls_runner_once(self):
        self.read(self.add()); calls=[]
        def model(c):calls.append(c);return fixture_model(c)
        with TestClient(create_app(self.path, core_only=True, materials_model_runner=model)) as client:
            revision=self.snap()['revision']
            a=self.generate(client,request_id='same-generation',expected_revision=revision)
            b=self.generate(client,request_id='same-generation',expected_revision=revision)
        self.assertEqual(a.json()['id'],b.json()['id']);self.assertEqual(len(calls),1)

    def test_partial_batch_excludes_failed_file_and_keeps_original(self):
        self.read(self.add());bad=self.add(base64.b64encode(b'not a pdf').decode(),kind='file',name='bad.pdf');self.read(bad)
        self.assertEqual(self.generate().status_code,200)
        s=self.snap();self.assertEqual(s['result']['scope']['excluded'][0]['id'],bad['id'])
        r=self.client.get(self.url+'/sources/'+bad['id']+'/original',headers=self.h('a'))
        self.assertEqual(r.content,b'not a pdf');self.assertEqual(s['sources'][-1]['metadata']['author'],'')

    def test_private_and_team_isolation_all_read_write_routes(self):
        s=self.prepared();source=s['sources'][0];item=s['result']['items'][0]
        for user in ('b','c'):
            for suffix in ('','/sources/'+source['id']+'/original'):
                self.assertEqual(self.client.get(self.url+suffix,headers=self.h(user)).status_code,403)
            self.assertEqual(self.client.post(self.url+'/sources/'+source['id']+'/parse',headers=self.h(user),json={'expected_revision':s['revision']}).status_code,403)
            self.assertEqual(self.client.patch(self.url+'/items/'+item['id'],headers=self.h(user),json={'text':'bad','verification':'unverified','note':'bad','expected_revision':s['revision']}).status_code,403)
        team=self.client.get('/api/materials/space?team_id='+self.sid,headers=self.h('a')).json()
        self.assertEqual(self.client.get('/api/materials/spaces/'+team['id'],headers=self.h('b')).status_code,200)
        self.assertEqual(self.client.get('/api/materials/spaces/'+team['id'],headers=self.h('c')).status_code,403)
        self.assertEqual(self.client.get(self.url).status_code,401)

    def test_bad_citation_member_and_statement_without_source_rejected(self):
        self.prepared();previous=copy.deepcopy(self.snap()['result'])
        for kind in ('quote','source','member','empty','html'):
            def bad(c):
                r=fixture_model(c);i=r['items'][0]
                if kind=='quote':i['citations'][0]['quote']='并不存在的原文'
                if kind=='source':i['citations'][0]['source_id']='other-team-source'
                if kind=='member':i['member_ids']=['c']
                if kind=='empty':i['citations']=[]
                if kind=='html':r['html']='<script>evil()</script>'
                return r
            with TestClient(create_app(self.path,core_only=True,materials_model_runner=bad)) as client:
                self.assertEqual(self.generate(client).status_code,502)
            self.assertEqual(self.snap()['result'],previous)

    def test_model_not_configured_and_timeout_keep_inputs(self):
        self.prepared();previous=self.snap()['result']
        with patch('app.materials.team_model_config',return_value={'api_key':''}),TestClient(create_app(self.path,core_only=True)) as client:
            r=self.generate(client);self.assertEqual(r.status_code,503);self.assertEqual(r.json()['error']['code'],'model_not_configured')
        def timeout(c):raise PlannerError('model_timeout','模型超时，原件保留。',504)
        with TestClient(create_app(self.path,core_only=True,materials_model_runner=timeout)) as client:
            self.assertEqual(self.generate(client).status_code,504)
        self.assertEqual(self.snap()['result'],previous)

    def test_edit_conflict_and_manual_protection(self):
        s=self.prepared();item=s['result']['items'][0]
        body={'expected_revision':s['revision'],'text':'人工更正：日期仍待核实','verification':'human_confirmed','note':'已对照原件'}
        url=self.url+'/items/'+item['id']
        self.assertEqual(self.client.patch(url,headers=self.h('a'),json=body).status_code,200)
        self.assertEqual(self.client.patch(url,headers=self.h('a'),json=body).status_code,409)
        self.read(self.add('新增：提交说明文档。',name='补充'))
        self.assertEqual(self.generate().status_code,200)
        updated=next(i for i in self.snap()['result']['items'] if i['id']==item['id'])
        self.assertEqual(updated['text'],body['text']);self.assertTrue(updated['protected'])
        self.assertTrue(any(h['kind']=='edit' for h in self.snap()['history']))

    def test_late_model_result_cannot_overwrite_new_source_or_edit(self):
        self.prepared()
        for change in ('source','edit'):
            started=threading.Event();release=threading.Event()
            def slow(c):started.set();release.wait(5);return fixture_model(c)
            with TestClient(create_app(self.path,core_only=True,materials_model_runner=slow)) as client,ThreadPoolExecutor() as pool:
                future=pool.submit(self.generate,client);self.assertTrue(started.wait(5))
                if change=='source':self.add('新增事实',name=self.key())
                else:
                    s=self.snap();i=s['result']['items'][0]
                    self.client.patch(self.url+'/items/'+i['id'],headers=self.h('a'),json={'expected_revision':s['revision'],'text':'受保护的修改','verification':'unverified','note':'人工修正'})
                release.set();self.assertEqual(future.result().status_code,409)
            self.assertEqual(self.snap()['batches'][0]['status'],'stale')

    def test_duplicate_facts_merge_citations_and_conflicts_remain(self):
        self.read(self.add('支持1至3人组队。',name='来源一'))
        self.read(self.add('支持1至3人组队。\n支持4至5人组队。',name='来源二'))
        self.assertEqual(self.generate().status_code,200)
        items=self.snap()['result']['items'];same=next(i for i in items if '1至3' in i['text'])
        self.assertEqual(len(same['citations']),2)
        self.assertEqual(len([i for i in items if i['conflict']]),2)

    def test_withdraw_invalidates_result_and_export(self):
        s=self.prepared();source=s['sources'][0]
        r=self.client.post(self.url+'/sources/'+source['id']+'/withdraw',headers=self.h('a'),json={'expected_revision':s['revision']})
        self.assertEqual(r.status_code,200);self.assertTrue(r.json()['stale'])
        self.assertEqual(r.json()['result']['items'][0]['verification'],'source_withdrawn')
        self.assertEqual(self.export().status_code,409)

    def export(self,**extra):
        s=self.snap();body={'request_id':self.key(),'expected_revision':s['revision'],'team_id':self.sid,'kind':'direction','share_sources':True,'item_ids':[s['result']['items'][0]['id']],**extra}
        return self.client.post(self.url+'/drafts',headers=self.h('a'),json=body)

    def test_shared_selection_is_a_draft_and_formal_permissions_unchanged(self):
        self.prepared();before=self.client.get(self.base,headers=self.h('a')).json()
        r=self.export(request_id='export-one');self.assertEqual(r.status_code,200,r.text)
        again=self.export(request_id='export-one');self.assertEqual(again.json()['id'],r.json()['id'])
        team=self.client.get('/api/materials/space?team_id='+self.sid,headers=self.h('b')).json()
        self.assertEqual(len(team['sources']),1);self.assertEqual(len(team['drafts']),1)
        after=self.client.get(self.base,headers=self.h('a')).json()
        self.assertEqual(before['direction'],after['direction']);self.assertEqual(before['active_plan'],after['active_plan'])
        self.assertEqual(self.client.put(self.base+'/direction',headers=self.h('b'),json={'goal':'bad','expected_revision':after['revision']}).status_code,403)
        self.assertEqual(self.export(team_id='other-team').status_code,403)

    def test_schema_restart_preserves_history_original_and_result(self):
        s=self.prepared()
        for _ in range(2):
            with TestClient(create_app(self.path,core_only=True)) as c:
                fresh=c.get(self.url,headers=self.h('a')).json()
                self.assertEqual(fresh['result'],s['result']);self.assertEqual(fresh['sources'],s['sources'])

    def test_same_team_draft_rejects_corrected_item_but_allows_unrelated_edits(self):
        self.space = self.client.get('/api/materials/space?team_id='+self.sid,headers=self.h('a')).json()['id']
        self.url = '/api/materials/spaces/'+self.space
        s = self.prepared();selected,other = s['result']['items'][:2]
        draft = self.export().json()
        def prepare():
            return self.client.post(self.url+'/drafts/'+draft['id']+'/prepare',headers=self.h('a'),json={'expected_revision':self.snap()['revision']})
        def edit(item):
            return self.client.patch(self.url+'/items/'+item['id'],headers=self.h('a'),json={'expected_revision':self.snap()['revision'],'text':'人工纠正：以核实后的内容为准','verification':'human_confirmed','note':'实际核对了原件'})
        self.assertEqual(prepare().status_code,200)
        self.assertEqual(edit(other).status_code,200)
        self.assertEqual(prepare().status_code,200)
        self.assertEqual(edit(selected).status_code,200)
        r=prepare();self.assertEqual(r.status_code,409);self.assertEqual(r.json()['error']['code'],'draft_items_changed')
        new=self.export().json()
        self.assertEqual(self.client.post(self.url+'/drafts/'+new['id']+'/prepare',headers=self.h('a'),json={'expected_revision':self.snap()['revision']}).status_code,200)

    def test_shared_copy_stays_usable_without_model_after_personal_edit(self):
        s=self.prepared();draft=self.export().json();item=s['result']['items'][0]
        self.client.patch(self.url+'/items/'+item['id'],headers=self.h('a'),json={'expected_revision':self.snap()['revision'],'text':'仅修改个人副本','verification':'human_confirmed','note':'不改变已分享副本'})
        own=self.client.post(self.url+'/drafts/'+draft['id']+'/prepare',headers=self.h('a'),json={'expected_revision':self.snap()['revision']})
        self.assertEqual(own.status_code,409)
        with TestClient(create_app(self.path,core_only=True)) as client:
            team=client.get('/api/materials/space?team_id='+self.sid,headers=self.h('b')).json()
            self.assertFalse(team['result'].get('items'))
            shared=team['drafts'][0]
            r=client.post('/api/materials/spaces/'+team['id']+'/drafts/'+shared['id']+'/prepare',headers=self.h('b'),json={'expected_revision':team['revision']})
            self.assertEqual(r.status_code,200,r.text)
            self.assertNotIn('仅修改个人副本',r.json()['body']['text'])

    def test_legacy_same_space_draft_compares_current_item(self):
        import sqlite3
        from contextlib import closing
        self.space=self.client.get('/api/materials/space?team_id='+self.sid,headers=self.h('a')).json()['id']
        self.url='/api/materials/spaces/'+self.space
        s=self.prepared();draft=self.export().json()
        with closing(sqlite3.connect(self.path)) as con:
            body=draft['body'];body.pop('item_basis')
            con.execute('UPDATE materials_drafts SET body=? WHERE id=?',(json.dumps(body),draft['id']))
            con.commit()
        url=self.url+'/drafts/'+draft['id']+'/prepare'
        self.assertEqual(self.client.post(url,headers=self.h('a'),json={'expected_revision':self.snap()['revision']}).status_code,200)
        item=s['result']['items'][0]
        self.client.patch(self.url+'/items/'+item['id'],headers=self.h('a'),json={'expected_revision':self.snap()['revision'],'text':'旧草稿对应条目已纠正','verification':'human_confirmed','note':'人工核对'})
        self.assertEqual(self.client.post(url,headers=self.h('a'),json={'expected_revision':self.snap()['revision']}).status_code,409)

    def test_real_txt_docx_and_pdf_parsing(self):
        from docx import Document
        from reportlab.pdfgen import canvas
        for encoding in ('utf-8-sig','gb18030'):
            segments,_,_=parse_source('file','a.txt','真实文本测试'.encode(encoding));self.assertIn('真实文本',segments[0]['text'])
        doc=Document();doc.add_paragraph('真实 DOCX 段落');doc.add_table(rows=1,cols=1).cell(0,0).text='表格文字';out=io.BytesIO();doc.save(out)
        segments,_,_=parse_source('file','a.docx',out.getvalue());self.assertEqual(len(segments),2)
        out=io.BytesIO();c=canvas.Canvas(out);c.drawString(50,700,'Submit a working demo.');c.save()
        segments,_,_=parse_source('file','a.pdf',out.getvalue());self.assertEqual(segments[0]['location'],'第 1 页');self.assertIn('working demo',segments[0]['text'])

    def test_ssrf_is_blocked_and_retains_url(self):
        for url in ('http://127.0.0.1:8893/','http://169.254.169.254/','http://localhost/','file:///etc/passwd'):
            s=self.read(self.add(url,kind='url',name=url))
            self.assertEqual(s['status'],'failed')
            r=self.client.get(self.url+'/sources/'+s['id']+'/original',headers=self.h('a'));self.assertEqual(r.text,url)

    def test_partial_pdf_failure_keeps_successful_pages(self):
        from reportlab.pdfgen import canvas
        out=io.BytesIO();c=canvas.Canvas(out);c.drawString(50,700,'Readable page');c.showPage();c.showPage();c.save()
        with patch('app.materials_parse.ocr',side_effect=RuntimeError('bad scan')):
            segments,warnings,_=parse_source('file','scan.pdf',out.getvalue())
        self.assertEqual(len(segments),1);self.assertIn('第 2 页未读入',warnings[0])

    def test_size_and_unsupported_formats(self):
        r=self.client.post(self.url+'/sources',headers=self.h('a'),json={'request_id':self.key(),'kind':'text','name':'大文件','content':'x'*40001})
        self.assertEqual(r.status_code,413)
        s=self.read(self.add(base64.b64encode(b'archive').decode(),kind='file',name='a.zip'))
        self.assertEqual(s['status'],'failed');self.assertIn('暂不支持',s['error'])

    def test_new_evidence_on_confirmed_item_requires_review(self):
        s=self.prepared();item=s['result']['items'][0]
        self.client.patch(self.url+'/items/'+item['id'],headers=self.h('a'),json={'expected_revision':s['revision'],'text':'我核对过的修正','verification':'human_confirmed','note':'核对原文'})
        self.read(self.add(item['text'],name='新的独立来源'))
        self.assertEqual(self.generate().status_code,200)
        current=next(i for i in self.snap()['result']['items'] if i['id']==item['id'])
        self.assertEqual(current['text'],'我核对过的修正');self.assertEqual(current['verification'],'needs_review')
        self.assertEqual(self.export(item_ids=[item['id']]).status_code,422)

    def test_shared_draft_fails_after_team_source_withdrawal(self):
        self.prepared();r=self.export();self.assertEqual(r.status_code,200)
        team=self.client.get('/api/materials/space?team_id='+self.sid,headers=self.h('a')).json();base='/api/materials/spaces/'+team['id']
        self.client.post(base+'/sources/'+team['sources'][0]['id']+'/withdraw',headers=self.h('b'),json={'expected_revision':team['revision']})
        r=self.client.post(self.url+'/drafts/'+r.json()['id']+'/prepare',headers=self.h('a'),json={'expected_revision':self.snap()['revision']})
        self.assertEqual(r.status_code,409)

    def test_export_shares_only_selected_sources_and_deduplicates_drafts(self):
        self.read(self.add('共享规则',name='共享资料'));self.read(self.add('不应共享的私人想法',name='私人材料'));self.generate()
        item=next(i for i in self.snap()['result']['items'] if i['text']=='共享规则')
        one=self.export(item_ids=[item['id']]);two=self.export(item_ids=[item['id']])
        self.assertEqual(one.json()['id'],two.json()['id'])
        team=self.client.get('/api/materials/space?team_id='+self.sid,headers=self.h('b')).json()
        self.assertEqual([s['metadata']['name'] for s in team['sources']],['共享资料'])

    def test_model_adapter_http_failure_and_valid_response(self):
        import httpx
        from app.materials import request_materials_model
        client_class=httpx.Client
        for status,code in ((401,'model_auth_failed'),(402,'model_insufficient_balance'),(429,'model_rate_limited'),(500,'model_unavailable')):
            transport=httpx.MockTransport(lambda r:httpx.Response(status,json={'error':'untrusted upstream body'}))
            with patch('app.materials.team_model_config',return_value={'api_key':'fake-key','base_url':'https://api.deepseek.com','model_id':'test'}),patch('app.materials.httpx.Client',side_effect=lambda **kw:client_class(transport=transport,**kw)):
                with self.assertRaises(PlannerError) as err:request_materials_model({})
                self.assertEqual(err.exception.code,code);self.assertNotIn('untrusted',err.exception.message)
        def response(request):
            payload=json.loads(request.content);self.assertEqual(len(payload['messages']),2);self.assertNotIn('tools',payload)
            return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':'{"summary":"ok","items":[]}'}}]})
        with patch('app.materials.team_model_config',return_value={'api_key':'fake-key','base_url':'https://api.deepseek.com','model_id':'test'}),patch('app.materials.httpx.Client',side_effect=lambda **kw:client_class(transport=httpx.MockTransport(response),**kw)):
            output,model=request_materials_model({'sources':[]});self.assertEqual(output['summary'],'ok');self.assertEqual(model,'test')

    def test_public_link_pin_dns_redirect_checks_and_html_text(self):
        with patch('app.contest_intel.socket.getaddrinfo', return_value=[(2, 1, 6, '', ('93.184.216.34', 443))]):
            import httpx
            from app.contest_intel import public_request, IntelError
            def redirect(request):return httpx.Response(302,headers={'location':'http://127.0.0.1/private'})
            with patch('app.contest_intel.socket.getaddrinfo',return_value=[(2,1,6,'',('93.184.216.34',443))]),httpx.Client(transport=httpx.MockTransport(redirect)) as client:
                with self.assertRaises(IntelError):public_request(client,'GET','https://example.com')
            response=httpx.Response(200,headers={'content-type':'text/html'},text='<html><body><p>真实正文</p><script>evil()</script></body></html>')
            with patch('app.materials_parse.public_request',return_value=(response,'https://example.com')):
                segments,warnings,original=parse_source('url','example',b'https://example.com')
                self.assertIn('真实正文',segments[0]['text']);self.assertNotIn('evil',segments[0]['text']);self.assertEqual(original,response.content)

    def test_late_parse_cannot_restore_withdrawn_material(self):
        source=self.add();started=threading.Event();release=threading.Event()
        original_parser=parse_source
        def slow(*args):started.set();release.wait(5);return original_parser(*args)
        with patch('app.materials.parse_source',side_effect=slow),ThreadPoolExecutor() as pool:
            future=pool.submit(self.client.post,self.url+'/sources/'+source['id']+'/parse',headers=self.h('a'),json={'expected_revision':self.snap()['revision']})
            self.assertTrue(started.wait(5))
            r=self.client.post(self.url+'/sources/'+source['id']+'/withdraw',headers=self.h('a'),json={'expected_revision':self.snap()['revision']});self.assertEqual(r.status_code,200)
            release.set();self.assertEqual(future.result().status_code,409)
        self.assertEqual(self.snap()['sources'][0]['status'],'withdrawn')


if __name__=='__main__':unittest.main()
