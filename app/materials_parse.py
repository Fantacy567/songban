"""Bounded local document parsing. OCR runs on Windows/Linux/macOS, without an API key."""
import io
import threading
import zipfile
from pathlib import Path

from .contest_intel import IntelError, html_to_text, public_request, validate_public_url
import httpx

MAX_BYTES = 8_000_000
MAX_CHARS = 40_000
MAX_PAGES = 20
_ocr = None
_ocr_lock = threading.Lock()


def ocr(raw):
    global _ocr
    try:
        from PIL import Image
        from rapidocr_onnxruntime import RapidOCR
        import numpy as np
    except ImportError:
        raise IntelError('ocr_unavailable', '本机未安装图片识别依赖，请安装 requirements-materials.txt，或补充文字。') from None
    with Image.open(io.BytesIO(raw)) as im:
        if im.width * im.height > 20_000_000:
            raise IntelError('image_too_large', '图片超过 2000 万像素，请缩小后重试。')
        im = im.convert('RGB')
        im.thumbnail((2400, 2400))
        pixels = np.asarray(im)[:, :, ::-1]
    with _ocr_lock:
        if _ocr is None:
            _ocr = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=2)
        result, _ = _ocr(pixels)
    lines = [line[1] for line in (result or []) if line[2] >= .5]
    if not lines:
        raise IntelError('ocr_empty', '没有识别到清晰文字，请补充文字或更清楚的图片。')
    return '\n'.join(lines)


def parse_source(kind, name, raw):
    segments, warnings = [], []
    original = raw

    def add(location, text, method='text'):
        # Keep long pages/text addressable without asking the model to copy huge passages.
        parts = []
        while len(text) > 1500:
            boundary = text.rfind('\n', 0, 1500)
            boundary = boundary + 1 if boundary > 500 else 1500
            parts.append(text[:boundary]); text = text[boundary:]
        parts.append(text)
        for n, part in enumerate(parts):
            if part.strip():
                label = location + (f' · 片段 {n + 1}' if len(parts) > 1 else '')
                segments.append({'id': str(len(segments) + 1), 'location': label, 'text': part.strip(), 'method': method})

    if kind == 'url':
        url = validate_public_url(raw.decode('utf-8'))
        try:
            with httpx.Client(timeout=20, follow_redirects=False, trust_env=False) as client:
                response, final = public_request(client, 'GET', url, pin_dns=True)
        except httpx.TimeoutException:
            raise IntelError('fetch_timeout', '公开网页连接或读取超时，链接已保留；可重试或粘贴正文。') from None
        except httpx.HTTPError:
            raise IntelError('fetch_failed', '无法连接公开网页，链接已保留；请核对地址或粘贴正文。') from None
        if response.status_code != 200:
            raise IntelError('fetch_failed', f'链接返回 HTTP {response.status_code}，可粘贴正文补充。')
        original = response.content
        if len(original) > MAX_BYTES:
            raise IntelError('web_too_large', '网页过大，未读取正文，请粘贴需要的段落。')
        content_type = response.headers.get('content-type', '').lower()
        if 'text/html' in content_type:
            text = html_to_text(response.text)
        elif 'text/plain' in content_type:
            text = response.text
        else:
            raise IntelError('web_format', '该链接没有返回公开网页正文，请下载文件后上传。')
        add('网页正文：' + final, text)
        warnings.append('仅读取无需登录的静态正文；动态内容、图片、登录墙未核验。')
    elif kind == 'text' or Path(name).suffix.lower() == '.txt':
        encodings = ('utf-16',) if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'gb18030')
        for encoding in encodings:
            try:
                text = raw.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            raise IntelError('encoding', '无法识别文本编码，请另存为 UTF-8。')
        add('正文', text)
    elif Path(name).suffix.lower() in ('.png', '.jpg', '.jpeg'):
        add('图片识别文字', ocr(raw), 'ocr')
        warnings.append('图片识别可能有误，请对照原件核实。')
    elif Path(name).suffix.lower() == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(raw))
        if pdf.is_encrypted:
            raise IntelError('encrypted_pdf', 'PDF 有密码保护，请提供可读取版本。')
        if len(pdf.pages) > MAX_PAGES:
            raise IntelError('page_limit', 'PDF 最多 20 页，请拆分后上传。')
        renderer = None
        try:
            for i, page in enumerate(pdf.pages):
                try:
                    text = page.extract_text() or ''
                    method = 'text'
                    if not text.strip():
                        import pypdfium2
                        if renderer is None:
                            renderer = pypdfium2.PdfDocument(raw)
                        rp = renderer[i]
                        try:
                            w, h = rp.get_size()
                            bitmap = rp.render(scale=min(2, 2400 / max(w, h)))
                            try:
                                im = bitmap.to_pil(); out = io.BytesIO(); im.save(out, format='PNG'); im.close()
                                text = ocr(out.getvalue()); method = 'ocr'
                            finally:
                                bitmap.close()
                        finally:
                            rp.close()
                    add(f'第 {i + 1} 页', text, method)
                    if method == 'ocr':
                        warnings.append(f'第 {i + 1} 页经过图片识别，可能有误，请对照原件核对。')
                except Exception as e:
                    warnings.append(f'第 {i + 1} 页未读入：' + getattr(e, 'message', '解析失败，请补充文字。'))
        finally:
            if renderer is not None:
                renderer.close()
    elif Path(name).suffix.lower() == '.docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40_000_000 or len(archive.infolist()) > 1000:
                raise IntelError('document_too_large', 'DOCX 解压后过大，请拆分文件。')
            if any(x.filename.startswith('word/media/') for x in archive.infolist()):
                warnings.append('DOCX 内嵌图片未识别，请单独上传图片；文本框、批注及复杂版式需人工核对。')
        doc = Document(io.BytesIO(raw))
        for i, paragraph in enumerate(doc.paragraphs):
            add(f'段落 {i + 1}', paragraph.text)
        for i, table in enumerate(doc.tables):
            for j, row in enumerate(table.rows):
                add(f'表 {i + 1} 行 {j + 1}', ' | '.join(c.text for c in row.cells))
    else:
        raise IntelError('unsupported_format', '暂不支持此格式。支持 PNG/JPEG、PDF、DOCX、TXT，可另补文字。')
    if not segments:
        raise IntelError('empty_source', '未读到可用文字。' + '；'.join(warnings))
    if sum(len(s['text']) for s in segments) > MAX_CHARS:
        raise IntelError('text_limit', '解析文字超过 4 万字，请拆分材料；未截断后冒充完整读取。')
    return segments, warnings, original
